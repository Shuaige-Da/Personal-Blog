"""Private, validated attachment handling for Hermes website chat."""

import base64
import io
import os
import uuid
from pathlib import Path

import filetype
from PIL import Image, UnidentifiedImageError
from docx import Document
from flask import current_app, url_for
from pypdf import PdfReader

from backend.core.models import HermesAttachment, db

IMAGE_MIMES = {'image/jpeg', 'image/png', 'image/webp', 'image/gif'}
IMAGE_EXTENSIONS = {'.jpg', '.jpeg', '.png', '.webp', '.gif'}
TEXT_EXTENSIONS = {
    '.txt', '.md', '.markdown', '.csv', '.json', '.py', '.js', '.ts', '.jsx',
    '.tsx', '.html', '.css', '.scss', '.xml', '.yaml', '.yml', '.toml', '.ini',
    '.cfg', '.sql', '.sh', '.ps1', '.java', '.c', '.h', '.cpp', '.cs', '.go',
    '.rs', '.log',
}
DOCUMENT_EXTENSIONS = TEXT_EXTENSIONS | {'.pdf', '.docx'}


class AttachmentError(ValueError):
    """One or more attachments did not pass validation."""


def attachment_root():
    root = os.path.abspath(current_app.config['HERMES_ATTACHMENT_FOLDER'])
    os.makedirs(root, exist_ok=True)
    return root


def attachment_path(attachment, *, extracted=False):
    name = attachment.extracted_text_name if extracted else attachment.stored_name
    if not name:
        return None
    root = attachment_root()
    path = os.path.abspath(os.path.join(root, name))
    if not path.startswith(root + os.sep):
        raise AttachmentError('附件路径无效。')
    return path


def _safe_original_name(raw_name):
    raw_name = os.path.basename((raw_name or '').replace('\\', '/')).strip()
    if not raw_name or len(raw_name) > 255:
        raise AttachmentError('附件文件名无效或过长。')
    return raw_name


def _decode_text(raw):
    for encoding in ('utf-8-sig', 'utf-8', 'gb18030'):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise AttachmentError('文本附件不是受支持的 UTF-8 或 GB18030 编码。')


def _extract_document(raw, extension):
    if extension in TEXT_EXTENSIONS:
        return _decode_text(raw)
    if extension == '.pdf':
        try:
            reader = PdfReader(io.BytesIO(raw))
            return '\n\n'.join(page.extract_text() or '' for page in reader.pages)
        except Exception as error:
            raise AttachmentError('PDF 无法安全解析。') from error
    if extension == '.docx':
        try:
            document = Document(io.BytesIO(raw))
            lines = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                lines.extend('\t'.join(cell.text for cell in row.cells) for row in table.rows)
            return '\n'.join(lines)
        except Exception as error:
            raise AttachmentError('DOCX 无法安全解析。') from error
    raise AttachmentError('不支持该文档类型。')


def _validate_image(raw, extension):
    detected = filetype.guess(raw[:4096])
    mime_type = detected.mime if detected else ''
    if extension not in IMAGE_EXTENSIONS or mime_type not in IMAGE_MIMES:
        raise AttachmentError('图片真实格式不受支持。')
    try:
        image = Image.open(io.BytesIO(raw))
        image.verify()
    except (UnidentifiedImageError, OSError, ValueError) as error:
        raise AttachmentError('图片无法安全解码。') from error
    return mime_type


def save_attachments(storages, *, conversation, message, user):
    """Validate and atomically save one turn's attachment collection."""
    storages = [storage for storage in storages if storage and storage.filename]
    max_count = current_app.config['HERMES_ATTACHMENT_MAX_COUNT']
    if len(storages) > max_count:
        raise AttachmentError(f'每条消息最多发送 {max_count} 个附件。')

    prepared = []
    total_size = 0
    for storage in storages:
        original_name = _safe_original_name(storage.filename)
        extension = Path(original_name).suffix.lower()
        raw = storage.read()
        size = len(raw)
        total_size += size
        is_image = extension in IMAGE_EXTENSIONS
        per_file_limit = current_app.config[
            'HERMES_ATTACHMENT_IMAGE_MAX_BYTES'
            if is_image else 'HERMES_ATTACHMENT_DOCUMENT_MAX_BYTES'
        ]
        if size <= 0 or size > per_file_limit:
            limit_mb = per_file_limit // 1024 // 1024
            raise AttachmentError(f'{original_name} 为空或超过 {limit_mb} MB 限制。')
        if total_size > current_app.config['HERMES_ATTACHMENT_TOTAL_MAX_BYTES']:
            total_limit_mb = (
                current_app.config['HERMES_ATTACHMENT_TOTAL_MAX_BYTES'] // 1024 // 1024
            )
            raise AttachmentError(f'本次附件总大小超过 {total_limit_mb} MB 限制。')

        if is_image:
            mime_type = _validate_image(raw, extension)
            extracted_text = None
            kind = 'image'
        else:
            if extension not in DOCUMENT_EXTENSIONS:
                raise AttachmentError(f'{original_name} 不是受支持的文档或代码文件。')
            extracted_text = _extract_document(raw, extension)
            mime_type = {
                '.pdf': 'application/pdf',
                '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            }.get(extension, 'text/plain')
            kind = 'document'
        prepared.append((original_name, extension, raw, mime_type, kind, extracted_text))

    root = attachment_root()
    saved_paths = []
    records = []
    try:
        for original_name, extension, raw, mime_type, kind, extracted_text in prepared:
            stored_name = f'{user.id}/{conversation.id}/{uuid.uuid4().hex}{extension}'
            absolute_path = os.path.join(root, *stored_name.split('/'))
            os.makedirs(os.path.dirname(absolute_path), exist_ok=True)
            with open(absolute_path, 'xb') as output:
                output.write(raw)
            saved_paths.append(absolute_path)

            extracted_name = None
            truncated = False
            if extracted_text is not None:
                max_chars = current_app.config['HERMES_ATTACHMENT_TEXT_MAX_CHARS']
                truncated = len(extracted_text) > max_chars
                extracted_text = extracted_text[:max_chars]
                extracted_name = stored_name + '.txt'
                extracted_path = os.path.join(root, *extracted_name.split('/'))
                with open(extracted_path, 'x', encoding='utf-8') as output:
                    output.write(extracted_text)
                saved_paths.append(extracted_path)

            record = HermesAttachment(
                conversation_id=conversation.id,
                message=message,
                user_id=user.id,
                stored_name=stored_name,
                original_name=original_name,
                mime_type=mime_type,
                size_bytes=len(raw),
                kind=kind,
                extracted_text_name=extracted_name,
                truncated=truncated,
            )
            db.session.add(record)
            records.append(record)
        db.session.flush()
        return records
    except Exception:
        for path in saved_paths:
            if os.path.isfile(path):
                os.remove(path)
        raise


def remove_attachment_files(attachments):
    directories = set()
    for attachment in attachments:
        for extracted in (False, True):
            path = attachment_path(attachment, extracted=extracted)
            if path and os.path.isfile(path):
                os.remove(path)
                directories.add(os.path.dirname(path))
    root = attachment_root()
    for directory in sorted(directories, key=len, reverse=True):
        while directory.startswith(root + os.sep) and directory != root:
            try:
                os.rmdir(directory)
            except OSError:
                break
            directory = os.path.dirname(directory)


def serialize_attachment(attachment):
    return {
        'id': attachment.id,
        'name': attachment.original_name,
        'mime_type': attachment.mime_type,
        'size_bytes': attachment.size_bytes,
        'kind': attachment.kind,
        'truncated': bool(attachment.truncated),
        'url': url_for('api_hermes_attachment', attachment_id=attachment.id),
    }


def build_attachment_content(message):
    """Build Hermes multimodal content parts from one saved user message."""
    parts = []
    document_sections = []
    total_chars = 0
    turn_limit = current_app.config['HERMES_ATTACHMENT_TURN_TEXT_MAX_CHARS']
    for attachment in message.attachments:
        if attachment.kind == 'image':
            path = attachment_path(attachment)
            with open(path, 'rb') as source:
                data = base64.b64encode(source.read()).decode('ascii')
            parts.append(
                {
                    'type': 'input_image',
                    'image_url': f'data:{attachment.mime_type};base64,{data}',
                    'detail': 'auto',
                }
            )
            continue
        extracted_path = attachment_path(attachment, extracted=True)
        if not extracted_path:
            continue
        with open(extracted_path, encoding='utf-8') as source:
            text = source.read()
        remaining = max(0, turn_limit - total_chars)
        included = text[:remaining]
        total_chars += len(included)
        marker = '（内容已截断）' if attachment.truncated or len(included) < len(text) else ''
        document_sections.append(
            f'\n\n[附件：{attachment.original_name}{marker}]\n{included}\n[/附件]'
        )
    text = (message.content or '').strip() + ''.join(document_sections)
    if text:
        parts.insert(0, {'type': 'input_text', 'text': text})
    return parts
